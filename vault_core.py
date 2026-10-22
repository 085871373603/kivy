"""Excel Vault data/storage layer, using pure Python and lightweight readers.

No pandas or NumPy are required. This keeps Android packaging simpler while
preserving CSV/XLSX/XLSM/XLS/ODS reading, indexing, searching and local storage.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import NamedTuple, Any

from openpyxl import load_workbook


def esc(value: str) -> str:
    """Escape user spreadsheet content for Kivy markup."""
    return str(value).replace("&", "&amp;").replace("[", "&bl;").replace("]", "&br;")


EXT = (".xlsx", ".xlsm", ".xls", ".ods", ".csv")
MAX_RESULTS = 3000
MAX_ROWS = 1500
MAX_ODS_REPEAT = 100000

# Desktop/Pydroid fallback. Android switches this to App.user_data_dir.
BASE = Path(__file__).resolve().parent / "storage"
FILES_DIR = BASE / "files"
PENDING_DIR = BASE / "pending"
INDEX = BASE / "index.json"
CACHE: dict[str, dict[str, "Sheet"]] = {}


class Sheet(NamedTuple):
    cols: list[str]
    rows: list[list[str]]  # rows after the header, each value stored as text
    first_row: int         # one-based spreadsheet row number of first data row


def configure_storage(root: str | Path, legacy_path: str | Path | None = None) -> None:
    """Select writable app-local storage and migrate the previous storage once."""
    global BASE, FILES_DIR, PENDING_DIR, INDEX
    new_base = Path(root).expanduser().resolve()
    old_base = Path(legacy_path).expanduser().resolve() if legacy_path else None
    BASE = new_base
    FILES_DIR = BASE / "files"
    PENDING_DIR = BASE / "pending"
    INDEX = BASE / "index.json"
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    PENDING_DIR.mkdir(parents=True, exist_ok=True)

    if old_base and old_base != BASE:
        old_index = old_base / "index.json"
        old_files = old_base / "files"
        if not INDEX.exists() and old_index.is_file():
            try:
                if old_files.is_dir():
                    shutil.copytree(old_files, FILES_DIR, dirs_exist_ok=True)
                shutil.copy2(old_index, INDEX)
            except OSError:
                # Migration is best-effort; fresh app storage still remains usable.
                pass
    CACHE.clear()


def load_index() -> list[dict]:
    try:
        data = json.loads(INDEX.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except (OSError, json.JSONDecodeError, TypeError):
        return []


def save_index(items: list[dict]) -> None:
    BASE.mkdir(parents=True, exist_ok=True)
    temporary = INDEX.with_suffix(".tmp")
    temporary.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, INDEX)


def col_letter(index: int) -> str:
    result, index = "", index + 1
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _guess_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines()[:30] if line.strip()]
    best, best_score = ",", 0
    for delimiter in ",;\t|":
        try:
            counts = [len(row) for row in csv.reader(lines, delimiter=delimiter)]
        except csv.Error:
            continue
        if counts and counts[0] > 1:
            score = counts.count(counts[0]) * 1000 + counts[0]
            if score > best_score:
                best, best_score = delimiter, score
    return best


def _stringify(value: Any) -> str:
    """Convert spreadsheet values to readable/searchable text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_csv(path: str | Path) -> list[list[str]]:
    raw_bytes = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw_bytes.decode("latin-1", errors="replace")
    try:
        csv.field_size_limit(2**30)
    except (OverflowError, ValueError):
        csv.field_size_limit(2**31 - 1)
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=_guess_delimiter(text)))
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows]


def _read_openpyxl(path: Path) -> dict[str, list[list[str]]]:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    output: dict[str, list[list[str]]] = {}
    try:
        for worksheet in workbook.worksheets:
            raw: list[list[str]] = []
            for values in worksheet.iter_rows(values_only=True):
                raw.append([_stringify(value) for value in values])
            output[str(worksheet.title)] = raw
    finally:
        workbook.close()
    return output


def _read_xls(path: Path) -> dict[str, list[list[str]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("Pembaca file .xls (xlrd) belum tersedia") from exc
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    output: dict[str, list[list[str]]] = {}
    try:
        for sheet in workbook.sheets():
            raw: list[list[str]] = []
            for row_index in range(sheet.nrows):
                values = []
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                        except (ValueError, OverflowError):
                            pass
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        value = ""
                    values.append(_stringify(value))
                raw.append(values)
            output[str(sheet.name)] = raw
    finally:
        workbook.release_resources()
    return output


def _read_ods(path: Path) -> dict[str, list[list[str]]]:
    try:
        from odf.opendocument import load as load_ods
        from odf.table import Table, TableRow
        from odf.teletype import extractText
    except ImportError as exc:
        raise RuntimeError("Pembaca file .ods (odfpy) belum tersedia") from exc

    document = load_ods(str(path))
    output: dict[str, list[list[str]]] = {}
    for table in document.spreadsheet.getElementsByType(Table):
        name = str(table.getAttribute("name") or "Sheet")
        raw: list[list[str]] = []
        for row_node in table.getElementsByType(TableRow):
            row_values: list[str] = []
            for cell in row_node.childNodes:
                if getattr(cell, "tagName", "") not in ("table:table-cell", "table:covered-table-cell"):
                    continue
                try:
                    repeat_columns = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                except (TypeError, ValueError):
                    repeat_columns = 1
                repeat_columns = max(1, min(repeat_columns, MAX_ODS_REPEAT))
                value = extractText(cell)
                if not value:
                    for attr in ("value", "datevalue", "timevalue", "booleanvalue", "stringvalue"):
                        candidate = cell.getAttribute(attr)
                        if candidate is not None:
                            value = str(candidate)
                            break
                row_values.extend([_stringify(value)] * repeat_columns)
            try:
                repeat_rows = int(row_node.getAttribute("numberrowsrepeated") or 1)
            except (TypeError, ValueError):
                repeat_rows = 1
            repeat_rows = max(1, min(repeat_rows, MAX_ODS_REPEAT - len(raw)))
            raw.extend([row_values.copy() for _ in range(repeat_rows)])
            if len(raw) >= MAX_ODS_REPEAT:
                break
        output[name] = raw
    return output


def _to_sheet(raw: list[list[Any]]) -> Sheet | None:
    """First non-empty row is the header; following rows become searchable data."""
    if not raw:
        return None
    normalized = [[_stringify(value) for value in row] for row in raw]
    width = max((len(row) for row in normalized), default=0)
    if width == 0:
        return None
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    header_index = None
    for index, row in enumerate(normalized):
        if any(value.strip() for value in row):
            header_index = index
            break
    if header_index is None:
        return None
    cols = [value.strip() or "Kolom " + col_letter(i) for i, value in enumerate(normalized[header_index])]
    rows = normalized[header_index + 1:]
    return Sheet(cols, rows, header_index + 2)


def read_sheets(path: str | Path) -> dict[str, Sheet]:
    """Read supported spreadsheet formats as text without pandas/NumPy."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        raw_sheets = {"CSV": _read_csv(file_path)}
    elif suffix in (".xlsx", ".xlsm"):
        raw_sheets = _read_openpyxl(file_path)
    elif suffix == ".xls":
        raw_sheets = _read_xls(file_path)
    elif suffix == ".ods":
        raw_sheets = _read_ods(file_path)
    else:
        raise ValueError("Format file tidak didukung")
    output: dict[str, Sheet] = {}
    for name, raw in raw_sheets.items():
        sheet = _to_sheet(raw)
        if sheet is not None:
            output[str(name)] = sheet
    if not output:
        raise ValueError("File tidak berisi data")
    return output


def get_sheets(item: dict) -> dict[str, Sheet]:
    key = item["stored"]
    if key not in CACHE:
        CACHE[key] = read_sheets(FILES_DIR / key)
    return CACHE[key]


def store_uploaded_file(source: str | Path) -> tuple[dict, dict[str, Sheet]]:
    """Validate and copy a spreadsheet into permanent app-local storage."""
    source_path = Path(source)
    if source_path.suffix.lower() not in EXT:
        raise ValueError("Format file tidak didukung")
    if not source_path.is_file():
        raise FileNotFoundError("File sumber tidak ditemukan")
    sheets = read_sheets(source_path)
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    stored = f"{time.time_ns()}_{source_path.name}"
    while (FILES_DIR / stored).exists():
        stored = f"{time.time_ns()}_{source_path.name}"
    destination = FILES_DIR / stored
    shutil.copy2(source_path, destination)
    item = {
        "name": source_path.name,
        "stored": stored,
        "sheets": len(sheets),
        "rows": int(sum(len(sheet.rows) for sheet in sheets.values())),
        "added": time.strftime("%Y-%m-%d %H:%M"),
    }
    try:
        index = load_index()
        index.append(item)
        save_index(index)
    except Exception:
        try:
            destination.unlink(missing_ok=True)
        finally:
            CACHE.pop(stored, None)
        raise
    CACHE[stored] = sheets
    return item, sheets


def clean(value: str) -> str:
    return str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")


def highlight(text: str, query: str) -> str:
    output, start = [], 0
    for match in re.finditer(re.escape(query), text, re.IGNORECASE):
        output.append(esc(text[start:match.start()]))
        output.append("[color=d9480f][b]%s[/b][/color]" % esc(match.group()))
        start = match.end()
    output.append(esc(text[start:]))
    return "".join(output)


def format_row(cols, values, hit=None, q=None, limit=80) -> str:
    parts = []
    for column_index, (name, value) in enumerate(zip(cols, values)):
        value = _stringify(value)
        if value == "" and column_index != hit:
            continue
        value = clean(value)[:limit]
        header = (
            "[b][color=3b49df]%s[/color][/b]" if column_index == hit
            else "[color=6b788c]%s[/color]"
        ) % esc(str(name)[:30])
        parts.append("%s %s" % (header, highlight(value, q) if q else esc(value)))
    return "    ".join(parts)


def search_tables(tables: list[dict], query: str, limit: int = MAX_RESULTS) -> dict:
    """Case-insensitive substring search across data cells, excluding headers."""
    hits: list[tuple[int, int, int]] = []
    cell_count = 0
    row_count = 0
    counter: Counter[str] = Counter()
    normalized_query = query.strip().casefold()
    if not normalized_query:
        return {"cells": 0, "rows": 0, "counter": counter, "hits": hits}

    for table_index, table in enumerate(tables):
        table_rows = table["rows"]
        matching_rows: set[int] = set()
        for row_index, row in enumerate(table_rows):
            for column_index, raw_value in enumerate(row):
                value = _stringify(raw_value)
                if normalized_query in value.casefold():
                    cell_count += 1
                    matching_rows.add(row_index)
                    counter[value.strip()] += 1
                    if len(hits) < max(0, limit):
                        hits.append((table_index, row_index, column_index))
        row_count += len(matching_rows)

    return {"cells": cell_count, "rows": row_count, "counter": counter, "hits": hits}


def hit_card(table: dict, row_index: int, column_index: int, query: str) -> str:
    row = table["rows"][row_index]
    excel_row = table["first_row"] + row_index
    return (
        "[b][color=3b49df]%s, %s[/color][/b]\n"
        "Sel [b]%s%d[/b] kolom %s:  %s\n"
        "Baris [b]%d[/b]:  %s"
    ) % (
        esc(table["file"]), esc(table["sheet"]), col_letter(column_index), excel_row,
        esc(table["cols"][column_index][:30]),
        highlight(clean(row[column_index])[:300], query),
        excel_row, format_row(table["cols"], row, column_index, query),
    )
