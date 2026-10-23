"""
Excel Vault - penyimpanan & pencarian data Excel / Spreadsheet (Kivy)

Instalasi : pip install kivy openpyxl xlrd        (xlrd hanya untuk .xls)
Jalankan  : python main.py
Pengujian : pip install pytest openpyxl && pytest -v test_main.py
Android   : lihat README.md (APK dibangun otomatis lewat GitHub Actions)

Alur
  1. Upload    : + Tambah File -> daftar antrean -> Simpan -> konfirmasi
  2. Dashboard : data tersimpan, pencarian teks di dalam sel (live). Setiap hasil
                 menampilkan SEL yang cocok beserta seluruh BARIS dari sel itu.
                 Tombol Clear untuk pencarian baru.

Baris pertama yang berisi data dianggap judul kolom (tidak ikut dicari).
Penyimpanan lokal: folder data aplikasi (App.user_data_dir)/storage. Di Android ini
penyimpanan privat aplikasi: tidak butuh izin, tetap ada setelah aplikasi ditutup.
"""
import csv
import io
import json
import os
import re
import shutil
import time
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
from typing import NamedTuple

from kivy.app import App
from kivy.clock import Clock
from kivy.core.window import Window
from kivy.lang import Builder
from kivy.metrics import dp
from kivy.properties import ListProperty
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.filechooser import FileChooserListView
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.screenmanager import Screen, ScreenManager
from kivy.utils import escape_markup as esc
from kivy.utils import platform

BASE = Path(__file__).resolve().parent / "storage"
FILES_DIR = BASE / "files"
INDEX = BASE / "index.json"
EXT = (".xlsx", ".xlsm", ".xls", ".ods", ".csv")
MAX_RESULTS = 3000   # batas sel hasil pencarian yang ditampilkan
MAX_ROWS = 1500      # batas baris pada tampilan awal dashboard
CACHE = {}           # nama file tersimpan -> {sheet: Sheet}

INK = (0.106, 0.141, 0.188, 1)
ACCENT = (0.231, 0.286, 0.875, 1)
DANGER = (0.78, 0.22, 0.27, 1)
SURFACE = (1, 1, 1, 1)


# ----------------------------------------------------------------- data layer
class Sheet(NamedTuple):
    cols: list          # nama kolom (dari baris judul)
    rows: list          # baris data: list of list[str], semuanya selebar kolom
    first_row: int      # nomor baris Excel dari baris data pertama


def init_storage(base):
    """Tetapkan folder penyimpanan lokal (di Android: penyimpanan privat aplikasi)."""
    global BASE, FILES_DIR, INDEX
    BASE = Path(base)
    FILES_DIR = BASE / "files"
    INDEX = BASE / "index.json"
    FILES_DIR.mkdir(parents=True, exist_ok=True)


def load_index():
    try:
        data = json.loads(INDEX.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except Exception:
        return []


def save_index(items):
    BASE.mkdir(parents=True, exist_ok=True)
    tmp = INDEX.with_suffix(".tmp")
    tmp.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, INDEX)  # atomik: indeks tidak rusak bila terhenti di tengah


def col_letter(i):
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def kind_of(path):
    """'xlsx' | 'xls' | 'ods' | 'csv' | None. Dari ekstensi; bila tak dikenali, dari isi file
    (nama file dari Google Drive/penyimpanan Android kadang tanpa ekstensi)."""
    p = str(path).lower()
    for ext in EXT:
        if p.endswith(ext):
            return {".xlsm": "xlsx"}.get(ext, ext[1:])
    try:
        with open(path, "rb") as f:
            head = f.read(4)
        if head == b"\xd0\xcf\x11\xe0":
            return "xls"
        if head[:2] == b"PK":
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
            if "xl/workbook.xml" in names:
                return "xlsx"
            if "content.xml" in names:
                return "ods"
    except Exception:
        pass
    return None


def _s(v):
    """Nilai sel -> teks."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def _grid_xlsx(path):
    import openpyxl
    with open(path, "rb") as f:  # lewat file object: openpyxl menolak nama tanpa ekstensi .xlsx
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        try:
            out = {}
            for ws in wb.worksheets:
                if hasattr(ws, "reset_dimensions"):
                    ws.reset_dimensions()  # sebagian file salah menulis ukuran sheet-nya
                out[ws.title] = [[_s(v) for v in row] for row in ws.iter_rows(values_only=True)]
            return out
        finally:
            wb.close()


def _xls_cell(xlrd, wb, cell):
    t, v = cell.ctype, cell.value
    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if t == xlrd.XL_CELL_DATE:
        try:
            return str(xlrd.xldate_as_datetime(v, wb.datemode))
        except Exception:
            return str(v)
    if t == xlrd.XL_CELL_BOOLEAN:
        return "True" if v else "False"
    if t == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(v, "#ERR")
    return _s(v)


def _grid_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path, on_demand=True)
    try:
        return {sh.name: [[_xls_cell(xlrd, wb, sh.cell(r, c)) for c in range(sh.ncols)]
                          for r in range(sh.nrows)] for sh in wb.sheets()}
    finally:
        wb.release_resources()


_NS = {"t": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
       "o": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
       "x": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}


def _q(p, n):
    return "{%s}%s" % (_NS[p], n)


def _ods_text(p):
    out = []

    def walk(e):
        if e.text:
            out.append(e.text)
        for ch in e:
            if ch.tag == _q("x", "s"):
                out.append(" " * int(ch.get(_q("x", "c")) or 1))
            elif ch.tag == _q("x", "tab"):
                out.append("\t")
            elif ch.tag == _q("x", "line-break"):
                out.append("\n")
            else:
                walk(ch)
            if ch.tail:
                out.append(ch.tail)
    walk(p)
    return "".join(out)


def _ods_cell(c):
    vt = c.get(_q("o", "value-type"))
    if vt in ("float", "percentage", "currency"):
        v = c.get(_q("o", "value")) or ""
        try:
            f = float(v)
            return str(int(f)) if f.is_integer() and abs(f) < 1e15 else v
        except ValueError:
            return v
    if vt == "date":
        return (c.get(_q("o", "date-value")) or "").replace("T", " ")
    if vt == "time":
        return c.get(_q("o", "time-value")) or ""
    if vt == "boolean":
        return "True" if c.get(_q("o", "boolean-value")) == "true" else "False"
    return "\n".join(_ods_text(p) for p in c.findall(_q("x", "p")))


def _grid_ods(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("content.xml"))
    out = {}
    for tb in root.iter(_q("t", "table")):
        rows, blank = [], 0
        for tr in tb.iter(_q("t", "table-row")):
            rep = min(int(tr.get(_q("t", "number-rows-repeated")) or 1), 100000)
            cells, pend = [], 0
            for c in tr:
                if c.tag not in (_q("t", "table-cell"), _q("t", "covered-table-cell")):
                    continue
                n = min(int(c.get(_q("t", "number-columns-repeated")) or 1), 16384)
                v = _ods_cell(c)
                if v == "":
                    pend += n  # sel kosong berulang: hanya dihitung bila disusul sel berisi
                else:
                    cells.extend([""] * pend + [v] * n)
                    pend = 0
            if not cells:
                blank += rep
            else:
                rows.extend([[] for _ in range(blank)])
                blank = 0
                rows.extend(list(cells) for _ in range(rep))
        out[tb.get(_q("t", "name")) or "Sheet"] = rows
    return out


def _guess_delimiter(text):
    lines = [x for x in text.splitlines()[:30] if x.strip()]
    best, score = ",", 0
    for d in ",;\t|":
        counts = [len(r) for r in csv.reader(lines, delimiter=d)]
        if counts and counts[0] > 1:
            s = counts.count(counts[0]) * 1000 + counts[0]
            if s > score:
                best, score = d, s
    return best


def _read_csv(path):
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            pass
    else:
        text = data.decode("latin-1")
    csv.field_size_limit(2 ** 30)
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=_guess_delimiter(text)))


def _to_sheet(grid):
    """Grid mentah (posisi = posisi di Excel) -> Sheet, atau None bila tidak ada isi."""
    while grid and not any(v.strip() for v in grid[-1]):
        grid.pop()  # buang baris kosong di akhir
    if not grid:
        return None
    h = next(i for i, r in enumerate(grid) if any(v.strip() for v in r))  # baris judul
    width = 0
    for r in grid[h:]:
        for i in range(len(r) - 1, width - 1, -1):
            if r[i].strip():
                width = i + 1
                break
    head = grid[h]
    cols = [(clean(head[i]).strip() if i < len(head) else "") or "Kolom " + col_letter(i)
            for i in range(width)]
    rows = [(r + [""] * (width - len(r)))[:width] for r in grid[h + 1:]]
    return Sheet(cols, rows, h + 2)


def read_sheets(path):
    """Baca semua sheet sebagai teks -> {nama_sheet: Sheet}. Error bila tak ada data."""
    kind = kind_of(path)
    if kind is None:
        raise ValueError("Format file tidak didukung")
    reader = {"xlsx": _grid_xlsx, "xls": _grid_xls, "ods": _grid_ods,
              "csv": lambda p: {"CSV": _read_csv(p)}}[kind]
    out = {}
    for name, grid in reader(str(path)).items():
        sheet = _to_sheet(grid)
        if sheet:
            out[str(name)] = sheet
    if not out:
        raise ValueError("File tidak berisi data")
    return out


def get_sheets(item):
    key = item["stored"]
    if key not in CACHE:
        CACHE[key] = read_sheets(FILES_DIR / key)
    return CACHE[key]


# ---------------------------------------------------------- pencarian & tampilan
def clean(v):
    return v.replace("\r", " ").replace("\n", " ").replace("\t", " ")


def highlight(text, q):
    """Tandai teks yang cocok dengan kata kunci (markup Kivy)."""
    out, i = [], 0
    for m in re.finditer(re.escape(q), text, re.IGNORECASE):
        out.append(esc(text[i:m.start()]))
        out.append("[color=d9480f][b]%s[/b][/color]" % esc(m.group()))
        i = m.end()
    out.append(esc(text[i:]))
    return "".join(out)


def format_row(cols, values, hit=None, q=None, limit=80):
    """Satu baris utuh: 'kolom nilai'. Kolom sel yang dicari diberi warna aksen."""
    parts = []
    for c, (name, v) in enumerate(zip(cols, values)):
        if v == "" and c != hit:
            continue
        v = clean(v)[:limit]
        head = ("[b][color=3b49df]%s[/color][/b]" if c == hit else "[color=6b788c]%s[/color]") % esc(name[:30])
        parts.append("%s %s" % (head, highlight(v, q) if q else esc(v)))
    return "    ".join(parts)


def search_tables(tables, q, limit=MAX_RESULTS):
    """Cari q (substring, tanpa membedakan huruf besar/kecil) di seluruh sel data."""
    ql = q.lower()
    hits, cells, rows, counter = [], 0, 0, Counter()
    for ti, t in enumerate(tables):
        for r, row in enumerate(t["rows"]):
            if ql not in "\x1f".join(row).lower():  # saring cepat per baris
                continue
            n = 0
            for c, v in enumerate(row):
                if v and ql in v.lower():
                    n += 1
                    counter[v.strip()] += 1
                    if len(hits) < limit:
                        hits.append((ti, r, c))
            cells += n
            rows += n > 0
    return {"cells": cells, "rows": rows, "counter": counter, "hits": hits}


def hit_card(t, r, c, q):
    """Kartu hasil: lokasi file/sheet, SEL yang cocok, dan seluruh BARIS dari sel itu."""
    row = t["rows"][r]
    xl = t["first_row"] + r
    return ("[b][color=3b49df]%s, %s[/color][/b]\n"
            "Sel [b]%s%d[/b] kolom %s:  %s\n"
            "Baris [b]%d[/b]:  %s") % (
        esc(t["file"]), esc(t["sheet"]), col_letter(c), xl, esc(t["cols"][c][:30]),
        highlight(clean(row[c])[:300], q), xl, format_row(t["cols"], row, c, q))


def go(screen):
    App.get_running_app().root.current = screen


def android_pick(on_files):
    """Buka pemilih file Android (Storage Access Framework): tanpa izin penyimpanan.
    File terpilih disalin ke penyimpanan privat aplikasi, lalu on_files(daftar_path,
    jumlah_gagal) dipanggil di thread Kivy."""
    from androidstorage4kivy import Chooser, SharedStorage
    ss = SharedStorage()

    def on_choose(uris):  # dipanggil dari thread Android -> pindahkan ke thread Kivy
        def work(_dt):
            paths, failed = [], 0
            for u in uris or []:
                try:
                    p = ss.copy_from_shared(u)
                    if p:
                        paths.append(p)
                    else:
                        failed += 1
                except Exception:
                    failed += 1
            on_files(paths, failed)
        Clock.schedule_once(work, 0)

    chooser = Chooser(on_choose)
    chooser.choose_content("*/*", multiple=True)
    return chooser


def popup(title, msg, buttons, height=dp(230)):
    """buttons = [(label, callback|None, primary?)]"""
    box = BoxLayout(orientation="vertical", padding=dp(14), spacing=dp(14))
    lbl = Label(text=msg, markup=True, color=INK, halign="center", valign="middle")
    lbl.bind(width=lambda w, x: setattr(w, "text_size", (x, None)))
    box.add_widget(lbl)
    row = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(8))
    pop = Popup(title=title, content=box, auto_dismiss=False, size_hint=(None, None),
                size=(min(dp(460), Window.width * 0.92), height),
                background="", background_color=SURFACE, title_color=INK,
                separator_color=ACCENT)
    for label, cb, primary in buttons:
        b = Button(text=label, bold=True, background_normal="", background_down="",
                   background_color=ACCENT if primary else (0.9, 0.92, 0.95, 1),
                   color=SURFACE if primary else INK)

        def _click(_, cb=cb):
            pop.dismiss()
            if cb:
                cb()
        b.bind(on_release=_click)
        row.add_widget(b)
    box.add_widget(row)
    pop.open()


# --------------------------------------------------------------------- screens
class UploadScreen(Screen):
    queue = ListProperty([])

    def open_chooser(self):
        if platform == "android":
            try:
                self._chooser = android_pick(self._after_pick)  # simpan referensi
            except Exception as e:
                popup("Penyimpanan tidak dapat dibuka",
                      "Pemilih file Android gagal dimuat:\n" + esc(str(e)[:150]),
                      [("OK", None, True)])
            return
        fc = FileChooserListView(
            path=str(Path.home()), multiselect=True,
            filters=[lambda folder, name: name.lower().endswith(EXT)])
        box = BoxLayout(orientation="vertical", spacing=dp(8), padding=dp(8))
        box.add_widget(fc)
        row = BoxLayout(size_hint_y=None, height=dp(44), spacing=dp(8))
        cancel = Button(text="Batal", background_normal="", color=INK,
                        background_color=(0.9, 0.92, 0.95, 1))
        add = Button(text="Tambah file terpilih", bold=True, background_normal="",
                     background_color=ACCENT)
        row.add_widget(cancel)
        row.add_widget(add)
        box.add_widget(row)
        pop = Popup(title="Pilih file Excel atau spreadsheet", content=box,
                    size_hint=(0.95, 0.9), background="", background_color=SURFACE,
                    title_color=INK, separator_color=ACCENT)
        cancel.bind(on_release=pop.dismiss)

        def _add(*_):
            pop.dismiss()
            self._after_pick(fc.selection)
        add.bind(on_release=_add)
        pop.open()

    def _after_pick(self, paths, failed=0):
        _added, ignored = self.add_files(paths)
        if ignored + failed:
            popup("Sebagian file dilewati",
                  "[b]%d[/b] file dilewati: format tidak didukung atau gagal dibaca.\n"
                  "Format yang didukung: xlsx, xls, ods, csv." % (ignored + failed),
                  [("OK", None, True)], height=dp(220))

    def add_files(self, paths):
        """Tambah ke antrean. Mengembalikan (jumlah_ditambah, jumlah_diabaikan)."""
        added = ignored = 0
        for p in paths:
            if not os.path.isfile(p) or kind_of(p) is None:
                ignored += 1
            elif p not in self.queue:
                self.queue.append(p)
                added += 1
        self.refresh()
        return added, ignored

    def remove(self, path):
        self.queue.remove(path)
        self.refresh()

    def refresh(self):
        box = self.ids.qbox
        box.clear_widgets()
        for p in self.queue:
            row = BoxLayout(size_hint_y=None, height=dp(48), spacing=dp(8))
            kb = os.path.getsize(p) / 1024
            lbl = Label(text="%s   [color=6b788c](%s KB)[/color]" % (esc(os.path.basename(p)), format(kb, ",.0f")),
                        markup=True, color=INK, halign="left", valign="middle",
                        shorten=True, shorten_from="center")
            lbl.bind(size=lambda w, s: setattr(w, "text_size", s))
            btn = Button(text="Hapus", size_hint_x=None, width=dp(80), bold=True,
                         background_normal="", background_color=DANGER)
            btn.bind(on_release=lambda _, p=p: self.remove(p))
            row.add_widget(lbl)
            row.add_widget(btn)
            box.add_widget(row)
        n = len(self.queue)
        self.ids.hint.text = (
            "%d file siap disimpan. Tekan Simpan untuk melanjutkan." % n if n else
            "Belum ada file. Tekan Tambah File untuk memilih berkas .xlsx, .xls, .ods atau .csv.")

    def ask_save(self):
        if not self.queue:
            return popup("Belum ada file", "Tambahkan minimal satu file terlebih dahulu.",
                         [("OK", None, True)], height=dp(190))
        popup("Konfirmasi penyimpanan",
              "Simpan [b]%d[/b] file ke penyimpanan?" % len(self.queue),
              [("Batal", None, False), ("Ya, simpan", self.save, True)], height=dp(190))

    def save(self):
        FILES_DIR.mkdir(parents=True, exist_ok=True)
        index, ok, errors = load_index(), 0, []
        for p in list(self.queue):
            name = os.path.basename(p)
            try:
                sheets = read_sheets(p)  # sekaligus validasi file
                stored = "%d_%s" % (time.time_ns(), name)
                if not name.lower().endswith(EXT):
                    stored += "." + kind_of(p)  # file tanpa ekstensi: ekstensi dari isinya
                shutil.copy2(p, FILES_DIR / stored)
                index.append({"name": name, "stored": stored, "sheets": len(sheets),
                              "rows": int(sum(len(s.rows) for s in sheets.values())),
                              "added": time.strftime("%Y-%m-%d %H:%M")})
                CACHE[stored] = sheets
                self.queue.remove(p)
                ok += 1
            except Exception as e:
                errors.append("%s: %s" % (name, str(e)[:100]))
        save_index(index)
        self.refresh()
        msg = "[b]%d[/b] file berhasil disimpan." % ok
        if errors:
            msg += "\n\nGagal dibaca:\n" + "\n".join(esc(x) for x in errors)
        popup("Hasil penyimpanan", msg,
              [("Upload lagi", None, False), ("Buka Dashboard", lambda: go("dashboard"), True)],
              height=dp(230 + 30 * len(errors)))


class DashboardScreen(Screen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.tables, self.failed, self.nfiles, self._ev = [], [], 0, None

    def reload(self):
        self.tables, self.failed, self.nfiles = [], [], 0
        for it in load_index():
            try:
                for name, s in get_sheets(it).items():
                    self.tables.append({"file": it["name"], "sheet": name, "cols": s.cols,
                                        "rows": s.rows, "first_row": s.first_row})
                self.nfiles += 1
            except Exception:
                self.failed.append(str(it.get("name", "?")))
        self.run_search()

    def on_query(self, _text):
        """Pencarian live dengan jeda singkat agar tetap lancar saat mengetik."""
        if self._ev:
            self._ev.cancel()
        self._ev = Clock.schedule_once(lambda dt: self.run_search(), 0.2)

    def clear(self):
        if self._ev:
            self._ev.cancel()
        self.ids.q.text = ""
        self.ids.q.focus = True
        self.run_search()

    def _note(self):
        if not self.failed:
            return ""
        return "\n[color=c73a45]%d file gagal dibaca: %s[/color]" % (
            len(self.failed), esc(", ".join(self.failed)[:120]))

    def run_search(self):
        q = self.ids.q.text.strip()
        if not self.tables:
            self.ids.summary.text = "Belum ada data. Buka menu Upload untuk menambahkan file." + self._note()
            self.ids.rv.data = []
        elif q:
            self.show_matches(q)
        else:
            self.show_all()

    def show_all(self):
        items, rows, cells = [], 0, 0
        for t in self.tables:
            rows += len(t["rows"])
            cells += len(t["rows"]) * len(t["cols"])
            for i, row in enumerate(t["rows"]):
                if len(items) >= MAX_ROWS:
                    break
                items.append({"text": "[b][color=3b49df]%s, %s[/color][/b]\nBaris [b]%d[/b]:  %s" % (
                    esc(t["file"]), esc(t["sheet"]), t["first_row"] + i, format_row(t["cols"], row))})
        note = ""
        if rows > MAX_ROWS:
            note = "\n[color=6b788c]Menampilkan %d baris pertama. Gunakan pencarian untuk menemukan data lain.[/color]" % MAX_ROWS
        self.ids.summary.text = "[b]%d[/b] file, [b]%d[/b] sheet, [b]%s[/b] baris data, [b]%s[/b] sel%s%s" % (
            self.nfiles, len(self.tables), format(rows, ","), format(cells, ","), note, self._note())
        self.ids.rv.data = items
        self.ids.rv.scroll_y = 1

    def show_matches(self, q):
        res = search_tables(self.tables, q)
        if not res["cells"]:
            self.ids.summary.text = "Tidak ada sel yang mengandung [b]%s[/b]." % esc(q) + self._note()
            self.ids.rv.data = []
            return
        top = ",  ".join("%s x%d" % (esc(clean(k)[:30]), n) for k, n in res["counter"].most_common(8))
        note = ("\n[color=6b788c]Hanya %d hasil pertama yang ditampilkan.[/color]" % MAX_RESULTS
                if res["cells"] > MAX_RESULTS else "")
        self.ids.summary.text = (
            "[b]%s[/b] sel cocok dengan [b]%s[/b] pada [b]%s[/b] baris, [b]%d[/b] nilai berbeda\n"
            "[color=6b788c]Terbanyak: %s[/color]%s%s" % (
                format(res["cells"], ","), esc(q), format(res["rows"], ","),
                len(res["counter"]), top, note, self._note()))
        self.ids.rv.data = [{"text": hit_card(self.tables[ti], r, c, q)} for ti, r, c in res["hits"]]
        self.ids.rv.scroll_y = 1


# -------------------------------------------------------------------------- UI
KV = """
#:import dp kivy.metrics.dp
#:import sp kivy.metrics.sp
#:set INK (0.106, 0.141, 0.188, 1)
#:set MUTED (0.42, 0.47, 0.55, 1)
#:set ACCENT (0.231, 0.286, 0.875, 1)

<Btn@Button>:
    background_normal: ''
    background_down: ''
    background_color: ACCENT
    color: 1, 1, 1, 1
    bold: True

<Quiet@Button>:
    background_normal: ''
    background_down: ''
    background_color: 1, 1, 1, 1
    color: ACCENT
    bold: True

<Title@Label>:
    color: INK
    bold: True
    font_size: sp(22)
    halign: 'left'
    valign: 'middle'
    text_size: self.size

<Info@Label>:
    color: MUTED
    markup: True
    halign: 'left'
    valign: 'top'
    size_hint_y: None
    text_size: self.width, None
    height: self.texture_size[1] + dp(8)

<Card@Label>:
    color: INK
    markup: True
    halign: 'left'
    valign: 'middle'
    size_hint_y: None
    text_size: self.width - dp(28), None
    height: self.texture_size[1] + dp(22)
    canvas.before:
        Color:
            rgba: 1, 1, 1, 1
        Rectangle:
            pos: self.x, self.y + dp(2)
            size: self.width, self.height - dp(4)

<UploadScreen>:
    name: 'upload'
    BoxLayout:
        orientation: 'vertical'
        padding: dp(16)
        spacing: dp(12)
        BoxLayout:
            size_hint_y: None
            height: dp(48)
            spacing: dp(8)
            Title:
                text: 'Upload file'
            Quiet:
                text: 'Dashboard'
                size_hint_x: None
                width: dp(120)
                on_release: app.root.current = 'dashboard'
        Info:
            id: hint
            text: ''
            color: INK
        ScrollView:
            GridLayout:
                id: qbox
                cols: 1
                size_hint_y: None
                height: self.minimum_height
                spacing: dp(6)
        BoxLayout:
            size_hint_y: None
            height: dp(48)
            spacing: dp(8)
            Quiet:
                text: '+ Tambah File'
                on_release: root.open_chooser()
            Btn:
                text: 'Simpan'
                on_release: root.ask_save()

<DashboardScreen>:
    name: 'dashboard'
    on_pre_enter: root.reload()
    BoxLayout:
        orientation: 'vertical'
        padding: dp(16)
        spacing: dp(10)
        BoxLayout:
            size_hint_y: None
            height: dp(48)
            spacing: dp(8)
            Title:
                text: 'Dashboard'
            Quiet:
                text: 'Upload'
                size_hint_x: None
                width: dp(100)
                on_release: app.root.current = 'upload'
        BoxLayout:
            size_hint_y: None
            height: dp(46)
            spacing: dp(8)
            TextInput:
                id: q
                hint_text: 'Cari teks di dalam sel...'
                multiline: False
                write_tab: False
                font_size: sp(16)
                padding: dp(12), dp(12)
                background_normal: ''
                background_active: ''
                background_color: 1, 1, 1, 1
                foreground_color: INK
                hint_text_color: MUTED
                cursor_color: ACCENT
                on_text: root.on_query(self.text)
            Btn:
                text: 'Clear'
                size_hint_x: None
                width: dp(90)
                on_release: root.clear()
        Info:
            id: summary
            text: ''
            color: INK
        RecycleView:
            id: rv
            viewclass: 'Card'
            bar_width: dp(6)
            RecycleBoxLayout:
                default_size: None, None
                default_size_hint: 1, None
                size_hint_y: None
                height: self.minimum_height
                orientation: 'vertical'
                spacing: dp(2)
"""


class ExcelVaultApp(App):
    title = "Excel Vault"

    def build(self):
        init_storage(Path(self.user_data_dir) / "storage")  # penyimpanan lokal aplikasi
        Window.clearcolor = (0.933, 0.945, 0.965, 1)
        if platform in ("win", "linux", "macosx"):
            Window.size = (1000, 680)
        Builder.load_string(KV)
        sm = ScreenManager()
        up = UploadScreen()
        sm.add_widget(up)
        sm.add_widget(DashboardScreen())
        up.refresh()
        Window.bind(on_keyboard=self._on_key)
        if load_index():
            sm.current = "dashboard"
        return sm

    def _on_key(self, _window, key, *_args):
        """Tombol Kembali Android: dari Upload kembali ke Dashboard, bukan keluar aplikasi."""
        if key == 27 and self.root.current == "upload" and load_index():
            self.root.current = "dashboard"
            return True
        return False

    def on_pause(self):
        return True  # wajib di Android: tanpa ini aplikasi ditutup saat pemilih file terbuka

    def on_resume(self):
        pass


if __name__ == "__main__":
    ExcelVaultApp().run()
