"""
Pengujian Excel Vault  ->  pytest -v test_main.py

Kivy diganti stub ringan supaya logika (pembaca file, pencarian, penyimpanan,
alur Upload/Dashboard) bisa diuji otomatis tanpa layar. Tampilan KV diperiksa
secara statis; render visual diuji manual (lihat daftar uji di README/balasan).
"""
import ast
import configparser
import datetime
import os
import random
import re
import string
import sys
import time
import types
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))


# ------------------------------------------------------------------ stub Kivy
def _install_kivy_stub():
    def mod(name, **kw):
        m = types.ModuleType(name)
        m.__dict__.update(kw)
        sys.modules[name] = m

    class W:
        def __init__(self, **kw):
            self.children = []
            self.__dict__.update(kw)
        def bind(self, **kw): pass
        def add_widget(self, w): self.children.append(w)
        def clear_widgets(self): self.children = []
        def open(self): pass
        def dismiss(self): pass

    class ListProperty:
        def __init__(self, d): self.d = d
        def __set_name__(self, o, n): self.n = "_" + n
        def __get__(self, o, t=None):
            return self if o is None else o.__dict__.setdefault(self.n, list(self.d))

    class Ev:
        def __init__(self, fn): self.fn, self.cancelled = fn, False
        def cancel(self): self.cancelled = True

    class Clock:
        events = []
        @staticmethod
        def schedule_once(fn, t=0):
            e = Ev(fn); Clock.events.append(e); return e

    class App:
        running = None
        @classmethod
        def get_running_app(cls): return cls.running

    def escape_markup(t):  # identik dengan kivy.utils.escape_markup
        return t.replace("&", "&amp;").replace("[", "&bl;").replace("]", "&br;")

    for n in ("kivy", "kivy.uix", "kivy.core"):
        mod(n); sys.modules[n].__path__ = []
    mod("kivy.app", App=App)
    mod("kivy.clock", Clock=Clock)
    mod("kivy.core.window", Window=SimpleNamespace(width=1000, size=(1000, 680), clearcolor=None))
    mod("kivy.lang", Builder=SimpleNamespace(load_string=lambda s: None))
    mod("kivy.metrics", dp=lambda x: x, sp=lambda x: x)
    mod("kivy.properties", ListProperty=ListProperty)
    mod("kivy.uix.boxlayout", BoxLayout=type("BoxLayout", (W,), {}))
    mod("kivy.uix.button", Button=type("Button", (W,), {}))
    mod("kivy.uix.filechooser", FileChooserListView=type("FC", (W,), {"selection": []}))
    mod("kivy.uix.label", Label=type("Label", (W,), {}))
    mod("kivy.uix.popup", Popup=type("Popup", (W,), {}))
    mod("kivy.uix.screenmanager", Screen=type("Screen", (W,), {}), ScreenManager=type("SM", (W,), {}))
    mod("kivy.utils", escape_markup=escape_markup, platform="linux")


_install_kivy_stub()
import main  # noqa: E402

SOURCE = Path(main.__file__).read_text(encoding="utf-8")
SAMPLE = [["Nama", "Kota", "Usia"],
          ["Budi Santoso", "Jakarta", 30],
          ["Ani Budiarti", "Bandung", 25],
          ["Citra", "Budi Kemuning", 41],
          ["Dewi", "Surabaya", None]]


# ------------------------------------------------------------------- helpers
def grid(rows, r0=1, c0=1):
    return {"%s%d" % (get_column_letter(c0 + j), r0 + i): v
            for i, row in enumerate(rows) for j, v in enumerate(row) if v is not None}


def make_xlsx(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for a, v in cells.items():
            ws[a] = v
    wb.save(path)
    return str(path)


def tables_of(path, name="f.xlsx"):
    return [{"file": name, "sheet": n, "cols": s.cols, "rows": s.rows, "first_row": s.first_row}
            for n, s in main.read_sheets(path).items()]


ODS_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')


def _c(v=None, rep=1):
    """Sel ODS: None = kosong; str/angka/bool; rep = number-columns-repeated."""
    from xml.sax.saxutils import escape
    r = ' table:number-columns-repeated="%d"' % rep if rep != 1 else ""
    if v is None:
        return "<table:table-cell%s/>" % r
    if isinstance(v, bool):
        return '<table:table-cell office:value-type="boolean" office:boolean-value="%s"%s/>' % (str(v).lower(), r)
    if isinstance(v, (int, float)):
        return '<table:table-cell office:value-type="float" office:value="%s"%s><text:p>%s</text:p></table:table-cell>' % (v, r, v)
    return '<table:table-cell office:value-type="string"%s><text:p>%s</text:p></table:table-cell>' % (r, escape(v))


def _r(cells, rep=1):
    r = ' table:number-rows-repeated="%d"' % rep if rep != 1 else ""
    return "<table:table-row%s>%s</table:table-row>" % (r, "".join(cells))


def ods_rows(rows):
    return [_r([_c(v) for v in row]) for row in rows]


def make_ods(path, sheets):
    """sheets = {nama: [xml_baris, ...]} -> berkas .ods mini (struktur seperti LibreOffice)."""
    body = "".join('<table:table table:name="%s">%s</table:table>' % (n, "".join(r)) for n, r in sheets.items())
    xml = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content %s><office:body>'
           '<office:spreadsheet>%s</office:spreadsheet></office:body></office:document-content>' % (ODS_NS, body))
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", xml)
    return str(path)


TAG = re.compile(r"\[(/?)(b|color)(=[0-9a-f]{6})?\]")


def valid_markup(s):
    """Markup Kivy sah: tak ada kurung liar / tag asing, tag seimbang, '&' ter-escape."""
    t = s.replace("&bl;", "").replace("&br;", "").replace("&amp;", "")
    if "&" in t:
        return False
    stack = []
    for tok in re.findall(r"\[[^\[\]]*\]|[\[\]]", t):
        m = TAG.fullmatch(tok)
        if not m:
            return False
        if m.group(1):
            if not stack or stack.pop() != m.group(2):
                return False
        else:
            stack.append(m.group(2))
    return not stack


def make_ids():
    box = lambda **kw: SimpleNamespace(children=[], clear_widgets=lambda: None, add_widget=lambda w: None, **kw)
    ids = SimpleNamespace(q=box(text="", focus=False), summary=box(text=""),
                          rv=box(data=[], scroll_y=0), hint=box(text=""))
    ids.qbox = SimpleNamespace(children=[])
    ids.qbox.clear_widgets = lambda: ids.qbox.children.clear()
    ids.qbox.add_widget = lambda w: ids.qbox.children.append(w)
    return ids


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "BASE", tmp_path / "storage")
    monkeypatch.setattr(main, "FILES_DIR", tmp_path / "storage" / "files")
    monkeypatch.setattr(main, "INDEX", tmp_path / "storage" / "index.json")
    main.CACHE.clear()
    pops = []
    monkeypatch.setattr(main, "popup", lambda title, msg, buttons, height=0: pops.append(
        SimpleNamespace(title=title, msg=msg, btn={b[0]: b[1] for b in buttons})))
    nav = SimpleNamespace(root=SimpleNamespace(current=None))
    sys.modules["kivy.app"].App.running = nav
    sys.modules["kivy.clock"].Clock.events.clear()
    return SimpleNamespace(tmp=tmp_path, pops=pops, nav=nav)


def dash():
    d = main.DashboardScreen(); d.ids = make_ids(); return d


def upl():
    u = main.UploadScreen(); u.ids = make_ids(); return u


@pytest.fixture
def sample(env):
    return make_xlsx(env.tmp / "sample.xlsx", {"Data": grid(SAMPLE)})


# ------------------------------------------------------- 1. pemeriksaan statis
class TestStatic:
    def test_sintaks(self):
        compile(SOURCE, "main.py", "exec")

    def test_kv_indentasi_rapi(self):
        for n, line in enumerate(main.KV.splitlines(), 1):
            assert "\t" not in line, "tab di baris KV %d" % n
            if line.strip():
                assert (len(line) - len(line.lstrip())) % 4 == 0, "indentasi baris KV %d" % n

    def test_id_kv_cocok_dengan_python(self):
        declared = set(re.findall(r"^\s+id:\s*(\w+)", main.KV, re.M))
        used = set(re.findall(r"self\.ids\.(\w+)", SOURCE))
        assert used <= declared, used - declared

    def test_metode_root_ada(self):
        for m in set(re.findall(r"root\.(\w+)\(", main.KV)):
            assert hasattr(main.UploadScreen, m) or hasattr(main.DashboardScreen, m), m

    def test_kelas_dan_widget_kv_dikenal(self):
        for cls in re.findall(r"^<(\w+)>:", main.KV, re.M):
            assert hasattr(main, cls), cls
        dynamic = set(re.findall(r"^<(\w+)@\w+>:", main.KV, re.M))
        builtin = {"BoxLayout", "ScrollView", "GridLayout", "TextInput", "RecycleView", "RecycleBoxLayout", "Color", "Rectangle"}
        for w in set(re.findall(r"^\s+([A-Z]\w+):\s*$", main.KV, re.M)):
            assert w in dynamic | builtin, w
        assert "Card" in dynamic  # viewclass RecycleView

    def test_target_navigasi_valid(self):
        screens = set(re.findall(r"^\s+name:\s*'(\w+)'", main.KV, re.M))
        for target in re.findall(r"app\.root\.current = '(\w+)'", main.KV):
            assert target in screens
        for target in re.findall(r'go\("(\w+)"\)', SOURCE):
            assert target in screens


# ------------------------------------------------------------- 2. pembaca file
class TestReader:
    def test_xlsx_dasar(self, sample):
        s = main.read_sheets(sample)["Data"]
        assert s.cols == ["Nama", "Kota", "Usia"] and s.first_row == 2 and len(s.rows) == 4
        assert s.rows[0] == ["Budi Santoso", "Jakarta", "30"]
        assert s.rows[3][2] == ""  # sel kosong -> teks kosong, bukan 'None'

    def test_data_tidak_mulai_di_A1(self, env):
        p = make_xlsx(env.tmp / "off.xlsx", {"S": grid(SAMPLE, r0=3, c0=3)})
        s = main.read_sheets(p)["S"]
        assert s.first_row == 4 and s.cols[2:] == ["Nama", "Kota", "Usia"]
        assert s.cols[0] == "Kolom A" and s.cols[1] == "Kolom B"

    def test_header_kosong_dan_duplikat(self, env):
        p = make_xlsx(env.tmp / "h.xlsx", {"S": {"A1": "Nama", "B1": None, "C1": "Nama", "A2": "x", "B2": "y", "C2": "z"}})
        s = main.read_sheets(p)["S"]
        assert s.cols == ["Nama", "Kolom B", "Nama"]
        assert main.search_tables(tables_of(p), "z")["cells"] == 1

    def test_banyak_sheet_dan_sheet_kosong_dilewati(self, env):
        p = make_xlsx(env.tmp / "m.xlsx", {"A": grid(SAMPLE), "Kosong": {}, "B": grid(SAMPLE[:2])})
        assert list(main.read_sheets(p)) == ["A", "B"]

    def test_file_kosong_ditolak(self, env):
        p = make_xlsx(env.tmp / "e.xlsx", {"S": {}})
        with pytest.raises(ValueError, match="tidak berisi data"):
            main.read_sheets(p)

    def test_hanya_judul_tanpa_data_tidak_error(self, env):
        p = make_xlsx(env.tmp / "ho.xlsx", {"S": {"A1": "Nama", "B1": "Kota"}})
        t = tables_of(p)
        assert not t[0]["rows"] and main.search_tables(t, "nama")["cells"] == 0

    def test_tipe_data(self, env):
        import datetime
        p = make_xlsx(env.tmp / "t.xlsx", {"S": {"A1": "h", "A2": 1000, "A3": 3.5, "A4": True,
                                                  "A5": datetime.datetime(2024, 1, 15), "A6": "teks"}})
        assert [r[0] for r in main.read_sheets(p)["S"].rows] == ["1000", "3.5", "True", "2024-01-15 00:00:00", "teks"]

    def test_ods_dasar(self, env):
        p = make_ods(env.tmp / "x.ods", {"Data": ods_rows(SAMPLE)})
        s = main.read_sheets(p)["Data"]
        assert s.cols == ["Nama", "Kota", "Usia"] and s.first_row == 2 and s.rows[0] == ["Budi Santoso", "Jakarta", "30"]
        assert main.search_tables(tables_of(p), "budi")["cells"] == 3

    @pytest.mark.parametrize("delim", [",", ";", "\t", "|"])
    def test_csv_delimiter(self, env, delim):
        p = env.tmp / "a.csv"
        p.write_text(delim.join(["Nama", "Kota"]) + "\n" + delim.join(["Budi", "Jakarta"]) + "\n", encoding="utf-8")
        s = main.read_sheets(p)["CSV"]
        assert s.cols == ["Nama", "Kota"] and s.rows[0] == ["Budi", "Jakarta"]

    def test_csv_satu_kolom_dan_baris_kosong_di_awal(self, env):
        p = env.tmp / "b.csv"
        p.write_text("\n\nNama\nBudi\n\nAni\n", encoding="utf-8")  # bug awal: 'Could not determine delimiter'
        s = main.read_sheets(p)["CSV"]
        assert s.cols == ["Nama"] and s.first_row == 4
        assert [r[0] for r in s.rows] == ["Budi", "", "Ani"]

    def test_csv_bom_cp1252_dan_baris_tak_rata(self, env):
        (env.tmp / "c.csv").write_bytes("\ufeffNama;Kota\nCaf\u00e9;Jakarta;ekstra\nBudi\n".encode("utf-8"))
        s = main.read_sheets(env.tmp / "c.csv")["CSV"]
        assert s.cols[:2] == ["Nama", "Kota"] and (len(s.rows), len(s.rows[0])) == (2, 3) and s.rows[0][0] == "Café"
        (env.tmp / "d.csv").write_bytes("Nama,Kota\nCaf\xe9,Bogor\n".encode("latin-1"))
        assert main.read_sheets(env.tmp / "d.csv")["CSV"].rows[0][0] == "Café"

    def test_csv_kutip_dan_baris_baru_dalam_sel(self, env):
        (env.tmp / "q.csv").write_text('Nama,Alamat\n"Budi","Jl. A,\nNo. 5"\n', encoding="utf-8")
        s = main.read_sheets(env.tmp / "q.csv")["CSV"]
        assert s.rows[0][1] == "Jl. A,\nNo. 5"

    def test_csv_kosong_dan_file_rusak(self, env):
        (env.tmp / "z.csv").write_text("", encoding="utf-8")
        with pytest.raises(ValueError):
            main.read_sheets(env.tmp / "z.csv")
        (env.tmp / "rusak.xlsx").write_bytes(b"ini bukan excel")
        with pytest.raises(Exception):
            main.read_sheets(env.tmp / "rusak.xlsx")


# --------------------------------------------------------------- 3. pencarian
class TestSearch:
    def test_hitungan_dan_urutan(self, sample):
        r = main.search_tables(tables_of(sample), "budi")  # Budi Santoso, Budiarti, Budi Kemuning
        assert (r["cells"], r["rows"], len(r["hits"])) == (3, 3, 3)
        assert [(h[1], h[2]) for h in r["hits"]] == [(0, 0), (1, 0), (2, 1)]  # urut baris lalu kolom
        assert set(r["counter"]) == {"Budi Santoso", "Ani Budiarti", "Budi Kemuning"}

    def test_tanpa_membedakan_huruf_dan_tanpa_hasil(self, sample):
        t = tables_of(sample)
        assert main.search_tables(t, "JAKARTA")["cells"] == 1
        assert main.search_tables(t, "tidak-ada")["cells"] == 0

    def test_beberapa_sel_dalam_satu_baris(self, env):
        p = make_xlsx(env.tmp / "r.xlsx", {"S": grid([["a", "b", "c"], ["kopi pahit", "kopi manis", "teh"]])})
        r = main.search_tables(tables_of(p), "kopi")
        assert r["cells"] == 2 and r["rows"] == 1

    def test_angka_dicari_sebagai_teks(self, sample):
        assert main.search_tables(tables_of(sample), "25")["cells"] == 1

    @pytest.mark.parametrize("q", ["[", "]", "&", "[b]", "[/color]", "%s", "%d", ".*", "(", "\\", "a|b", "{0}", "é", "İ", "日本"])
    def test_karakter_khusus_literal_dan_markup_valid(self, env, q):
        cells = grid([["h", "i"], ["x [b]tebal[/b] & 100% (a|b) .* \\ é İ 日本 %s %d {0}", "plain"]])
        t = tables_of(make_xlsx(env.tmp / "s.xlsx", {"S": cells}))
        r = main.search_tables(t, q)
        for ti, row, c in r["hits"]:
            assert valid_markup(main.hit_card(t[ti], row, c, q)), q

    def test_batas_hasil_tapi_hitungan_tetap_akurat(self):
        t = [{"file": "f", "sheet": "s", "cols": ["a", "b"], "rows": [["kopi", "teh"] for _ in range(100)], "first_row": 2}]
        r = main.search_tables(t, "kopi", limit=10)
        assert r["cells"] == 100 and r["rows"] == 100 and len(r["hits"]) == 10

    def test_alamat_sel_sama_dengan_file_excel_asli(self, env):
        """Uji properti: untuk SETIAP hit, sel di file Excel asli memang berisi kata kunci."""
        rng = random.Random(7)
        words = ["apel", "Jeruk", "mangga", "budi", "BUDI x", "kota a", "x"]
        cells = {"%s3" % get_column_letter(c): "H%d" % c for c in range(2, 9)}  # judul di baris 3
        for r in range(4, 80):
            for c in range(2, 9):
                if rng.random() < 0.6:
                    cells["%s%d" % (get_column_letter(c), r)] = rng.choice(words) + str(rng.randint(0, 9))
        p = make_xlsx(env.tmp / "p.xlsx", {"S": cells})
        t = tables_of(p)
        ws = openpyxl.load_workbook(p, data_only=True)["S"]
        for q in ["budi", "apel1", "x", "KOTA A", "9"]:
            r = main.search_tables(t, q)
            expected = sum(1 for row in ws.iter_rows(min_row=4) for c in row
                           if c.value is not None and q.lower() in str(c.value).lower())
            assert r["cells"] == expected, q
            for ti, rr, cc in r["hits"]:
                addr = "%s%d" % (main.col_letter(cc), t[ti]["first_row"] + rr)
                assert q.lower() in str(ws[addr].value).lower(), (q, addr)
                assert ("Sel [b]%s[/b]" % addr) in main.hit_card(t[ti], rr, cc, q)

    def test_fuzz_markup_selalu_valid(self, env):
        rng = random.Random(11)
        alpha = list("[]&%{}/=; \n\t") + ["[b]", "[/b]", "[color=ff0000]", "&amp;", "&bl;", "abc", "ABC", "é"]
        junk = lambda: "".join(rng.choice(alpha) for _ in range(rng.randint(0, 12)))
        rows = [[junk() for _ in range(5)] for _ in range(60)]
        t = [{"file": junk() or "f", "sheet": junk() or "s", "cols": [junk() or "k" for _ in range(5)], "rows": rows, "first_row": 2}]
        n = 0
        for q in ["a", "[", "&", "[b]", "&amp;", "abc", "é", "%", "/b"]:
            for ti, r, c in main.search_tables(t, q)["hits"]:
                assert valid_markup(main.hit_card(t[ti], r, c, q)), (q, r, c)
                n += 1
        assert n > 100


# ------------------------------------- 4. koreksi: hasil = SEL + BARIS dari sel
class TestCellAndRowDisplay:
    def test_kartu_menampilkan_sel_dan_seluruh_baris(self, env, sample):
        d = dash()
        main.save_index([{"name": "sample.xlsx", "stored": "k1.xlsx", "sheets": 1, "rows": 4, "added": ""}])
        main.FILES_DIR.mkdir(parents=True)
        Path(sample).replace(main.FILES_DIR / "k1.xlsx")
        d.reload()
        d.ids.q.text = "jakarta"
        d.run_search()
        card = d.ids.rv.data[0]["text"]
        assert len(d.ids.rv.data) == 1
        assert "Sel [b]B2[/b] kolom Kota" in card                     # SEL yang cocok
        assert "Baris [b]2[/b]:" in card                              # BARIS dari sel itu
        for isi in ("Budi Santoso", "Usia", "30", "Nama"):            # isi seluruh baris tampil
            assert isi in card
        assert "[color=d9480f][b]Jakarta[/b][/color]" in card         # kata kunci disorot
        assert "[b][color=3b49df]Kota[/color][/b]" in card            # kolom sel ditandai di baris
        assert "Bandung" not in card and "Ani" not in card            # baris lain tidak ikut
        assert "1" in d.ids.summary.text and "sel cocok" in d.ids.summary.text

    def test_dua_sel_dalam_satu_baris_masing_masing_punya_kartu(self, env):
        p = make_xlsx(env.tmp / "r.xlsx", {"S": grid([["a", "b", "c"], ["kopi pahit", "kopi manis", "teh"]])})
        t = tables_of(p)
        cards = [main.hit_card(t[0], r, c, "kopi") for _, r, c in main.search_tables(t, "kopi")["hits"]]
        assert "Sel [b]A2[/b]" in cards[0] and "Sel [b]B2[/b]" in cards[1]
        assert all("Baris [b]2[/b]" in c and "teh" in c for c in cards)

    def test_nilai_dengan_baris_baru_ditampilkan_satu_baris(self, env):
        p = make_xlsx(env.tmp / "n.xlsx", {"S": grid([["h", "i"], ["Ani\nSari", "x"]])})
        t = tables_of(p)
        card = main.hit_card(t[0], 0, 0, "sari")
        assert card.count("\n") == 2  # tepat 3 baris tampilan: lokasi, sel, baris


# ------------------------------------------------------------ 5. alur Upload
class TestUploadFlow:
    def test_tambah_file_validasi_dan_dedupe(self, env, sample):
        u = upl()
        (env.tmp / "catatan.txt").write_text("x")
        u.add_files([sample, sample, str(env.tmp / "catatan.txt"), str(env.tmp / "hilang.xlsx")])
        assert u.queue == [sample] and len(u.ids.qbox.children) == 1
        assert "1 file siap" in u.ids.hint.text
        u.remove(sample)
        assert u.queue == [] and "Belum ada file" in u.ids.hint.text

    def test_simpan_tanpa_file_diberi_peringatan(self, env):
        upl().ask_save()
        assert env.pops[-1].title == "Belum ada file"

    def test_batal_tidak_menyimpan_lalu_konfirmasi_menyimpan(self, env, sample):
        u = upl(); u.add_files([sample]); u.ask_save()
        konfirmasi = env.pops[-1]
        assert konfirmasi.title == "Konfirmasi penyimpanan" and "1" in konfirmasi.msg
        assert konfirmasi.btn["Batal"] is None
        assert not main.INDEX.exists()                                   # belum ada yang tersimpan
        konfirmasi.btn["Ya, simpan"]()
        idx = main.load_index()
        assert len(idx) == 1 and idx[0]["name"] == "sample.xlsx" and idx[0]["rows"] == 4
        assert (main.FILES_DIR / idx[0]["stored"]).exists() and u.queue == []
        hasil = env.pops[-1]
        assert "1" in hasil.msg and "berhasil" in hasil.msg
        hasil.btn["Buka Dashboard"]()
        assert env.nav.root.current == "dashboard"
        assert not list(main.BASE.glob("*.tmp"))                         # tak ada sisa file sementara

    def test_sebagian_gagal_file_baik_tetap_tersimpan(self, env, sample):
        rusak = env.tmp / "rusak.xlsx"; rusak.write_bytes(b"bukan excel")
        kosong = make_xlsx(env.tmp / "kosong.xlsx", {"S": {}})
        u = upl(); u.add_files([str(rusak), sample, kosong]); u.ask_save()
        env.pops[-1].btn["Ya, simpan"]()
        assert [i["name"] for i in main.load_index()] == ["sample.xlsx"]
        assert "rusak.xlsx" in env.pops[-1].msg and "kosong.xlsx" in env.pops[-1].msg
        assert sorted(os.path.basename(p) for p in u.queue) == ["kosong.xlsx", "rusak.xlsx"]  # tetap di antrean
        assert valid_markup(env.pops[-1].msg)

    def test_nama_file_berkarakter_markup_aman(self, env):
        p = make_xlsx(env.tmp / "a[b]&c.xlsx", {"S": grid(SAMPLE)})
        u = upl(); u.add_files([p]); u.ask_save(); env.pops[-1].btn["Ya, simpan"]()
        assert main.load_index()[0]["name"] == "a[b]&c.xlsx"

    def test_file_sama_disimpan_dua_kali_tidak_menimpa(self, env, sample):
        for _ in range(2):
            u = upl(); u.add_files([sample]); u.ask_save(); env.pops[-1].btn["Ya, simpan"]()
        idx = main.load_index()
        assert len(idx) == 2 and idx[0]["stored"] != idx[1]["stored"]
        assert len(list(main.FILES_DIR.iterdir())) == 2

    def test_semua_format_didukung_bisa_disimpan(self, env):
        a = make_xlsx(env.tmp / "a.xlsx", {"S": grid(SAMPLE)})
        b = make_ods(env.tmp / "b.ods", {"S": ods_rows(SAMPLE)})
        c = env.tmp / "c.csv"; c.write_text("Nama,Kota\nBudi,Jakarta\n", encoding="utf-8")
        u = upl(); u.add_files([a, b, str(c)]); u.ask_save(); env.pops[-1].btn["Ya, simpan"]()
        assert len(main.load_index()) == 3 and u.queue == []


# ---------------------------------------------------------- 6. alur Dashboard
class TestDashboard:
    def seed(self, env, *paths):
        u = upl(); u.add_files(list(paths)); u.ask_save(); env.pops[-1].btn["Ya, simpan"]()

    def test_keadaan_kosong(self, env):
        d = dash(); d.reload()
        assert "Belum ada data" in d.ids.summary.text and d.ids.rv.data == []

    def test_data_tersimpan_tampil_dan_bertahan_setelah_restart(self, env, sample):
        self.seed(env, sample)
        main.CACHE.clear()                       # simulasi aplikasi dibuka ulang
        d = dash(); d.reload()
        assert len(d.ids.rv.data) == 4 and "4" in d.ids.summary.text and "12" in d.ids.summary.text
        assert "Baris [b]2[/b]" in d.ids.rv.data[0]["text"] and "Budi Santoso" in d.ids.rv.data[0]["text"]

    def test_pencarian_live_debounce_dan_clear(self, env, sample):
        self.seed(env, sample)
        d = dash(); d.reload()
        Clock = sys.modules["kivy.clock"].Clock
        d.ids.q.text = "b"; d.on_query("b")
        d.ids.q.text = "budi"; d.on_query("budi")
        e1, e2 = Clock.events[-2:]
        assert e1.cancelled and not e2.cancelled                         # ketikan lama dibatalkan
        e2.fn(0)
        assert len(d.ids.rv.data) == 3 and "3" in d.ids.summary.text.split("sel cocok")[0]
        d.clear()                                                        # tombol Clear
        assert d.ids.q.text == "" and d.ids.q.focus is True
        assert len(d.ids.rv.data) == 4 and "baris data" in d.ids.summary.text
        d.ids.q.text = "   "; d.run_search()                             # spasi saja = kosong
        assert len(d.ids.rv.data) == 4

    def test_tanpa_hasil(self, env, sample):
        self.seed(env, sample)
        d = dash(); d.reload(); d.ids.q.text = "zzz"; d.run_search()
        assert "Tidak ada sel" in d.ids.summary.text and d.ids.rv.data == []

    def test_pencarian_lintas_file_dan_sheet(self, env):
        a = make_xlsx(env.tmp / "a.xlsx", {"S1": grid(SAMPLE), "S2": grid([["x"], ["Budi lain"]])})
        c = env.tmp / "c.csv"; c.write_text("Nama\nBudi csv\n", encoding="utf-8")
        self.seed(env, a, str(c))
        d = dash(); d.reload(); d.ids.q.text = "budi"; d.run_search()
        txt = " ".join(x["text"] for x in d.ids.rv.data)
        assert len(d.ids.rv.data) == 5 and "a.xlsx, S2" in txt and "c.csv, CSV" in txt
        assert "5" in d.ids.summary.text and "nilai berbeda" in d.ids.summary.text

    def test_batas_tampilan_awal(self, env, sample, monkeypatch):
        self.seed(env, sample)
        monkeypatch.setattr(main, "MAX_ROWS", 3)
        d = dash(); d.reload()
        assert len(d.ids.rv.data) == 3 and "3 baris pertama" in d.ids.summary.text

    def test_file_tersimpan_hilang_dilaporkan_bukan_crash(self, env, sample):
        self.seed(env, sample)
        idx = main.load_index() + [{"name": "hilang[1].xlsx", "stored": "tidak-ada", "sheets": 1, "rows": 1, "added": ""}]
        main.save_index(idx)
        d = dash(); d.reload()
        assert len(d.ids.rv.data) == 4 and "1 file gagal dibaca" in d.ids.summary.text
        assert valid_markup(d.ids.summary.text)

    def test_indeks_rusak_tidak_membuat_crash(self, env):
        main.BASE.mkdir(parents=True); main.INDEX.write_text("{rusak", encoding="utf-8")
        assert main.load_index() == []
        main.INDEX.write_text('{"a": 1}', encoding="utf-8")
        assert main.load_index() == []
        d = dash(); d.reload()
        assert "Belum ada data" in d.ids.summary.text

    def test_ringkasan_selalu_markup_valid(self, env):
        cells = grid([["[b]h&", "i]"], ["x [b] &amp; [", "[/color] y"]])
        p = make_xlsx(env.tmp / "w[1]&.xlsx", {"S&1": cells})
        self.seed(env, p)
        d = dash(); d.reload()
        for q in ["", "[", "&", "[b]", "zzz", "x"]:
            d.ids.q.text = q; d.run_search()
            assert valid_markup(d.ids.summary.text), q
            assert all(valid_markup(i["text"]) for i in d.ids.rv.data), q


# ----------------------------------------- 7. ODS, XLS, deteksi format, dimensi salah
class TestFormats:
    def test_ods_baris_kolom_berulang_dan_posisi(self, env):
        rows = [_r([_c(None, 1024)], rep=2),                                  # baris 1-2 kosong
                _r([_c(None, 2), _c("Nama"), _c("Kota")]),                    # judul di C3
                _r([_c(None, 2), _c("Budi"), _c("Jakarta")]),
                _r([_c(None, 1024)], rep=3),                                  # baris 5-7 kosong
                _r([_c(None, 2), _c("Ani"), _c("Bogor")]),                    # baris 8
                _r([_c(None, 1024)], rep=1048570)]                            # sisa kosong ala LibreOffice
        p = make_ods(env.tmp / "r.ods", {"S": rows})
        s = main.read_sheets(p)["S"]
        assert s.first_row == 4 and s.cols == ["Kolom A", "Kolom B", "Nama", "Kota"]
        assert len(s.rows) == 5 and s.rows[0] == ["", "", "Budi", "Jakarta"] and s.rows[4] == ["", "", "Ani", "Bogor"]
        t = tables_of(p); hit = main.search_tables(t, "ani")["hits"][0]
        assert "%s%d" % (main.col_letter(hit[2]), t[0]["first_row"] + hit[1]) == "C8"

    def test_ods_teks_spasi_paragraf_dan_entitas(self, env):
        raw = ('<table:table-cell office:value-type="string"><text:p>a<text:s text:c="2"/>b</text:p>'
               '<text:p>baris dua</text:p></table:table-cell>')
        p = make_ods(env.tmp / "t.ods", {"S": [_r([_c("h")]), _r([raw]), _r([_c("A & B <x>")])]})
        assert [r[0] for r in main.read_sheets(p)["S"].rows] == ["a  b\nbaris dua", "A & B <x>"]

    def test_ods_tipe_nilai(self, env):
        dt = ('<table:table-cell office:value-type="date" office:date-value="2024-01-15T10:30:00">'
              '<text:p>15/01/2024</text:p></table:table-cell>')
        pct = '<table:table-cell office:value-type="percentage" office:value="0.25"><text:p>25%</text:p></table:table-cell>'
        p = make_ods(env.tmp / "v.ods", {"S": [_r([_c("h")]), _r([_c(30)]), _r([_c(3.5)]), _r([_c(True)]), _r([dt]), _r([pct])]})
        assert [r[0] for r in main.read_sheets(p)["S"].rows] == ["30", "3.5", "True", "2024-01-15 10:30:00", "0.25"]

    def test_xls_dengan_xlrd_tiruan(self, env):
        """xlrd asli tidak ada di sandbox: konversi sel diuji dengan xlrd tiruan berantarmuka sama."""
        x = types.ModuleType("xlrd")
        (x.XL_CELL_EMPTY, x.XL_CELL_TEXT, x.XL_CELL_NUMBER, x.XL_CELL_DATE,
         x.XL_CELL_BOOLEAN, x.XL_CELL_ERROR, x.XL_CELL_BLANK) = range(7)
        x.error_text_from_code = {7: "#DIV/0!"}
        x.xldate_as_datetime = lambda v, mode: datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v)
        data = [[(1, "Nama"), (1, "Usia"), (1, "Lahir"), (1, "Aktif"), (1, "Err")],
                [(1, "Budi"), (2, 30.0), (3, 45306.0), (4, 1), (5, 7)],
                [(0, ""), (6, ""), (2, 3.5), (4, 0), (1, "x")]]
        sheet = SimpleNamespace(name="Lembar1", nrows=3, ncols=5,
                                cell=lambda r, c: SimpleNamespace(ctype=data[r][c][0], value=data[r][c][1]))
        wb = SimpleNamespace(datemode=0, sheets=lambda: [sheet], release_resources=lambda: None)
        x.open_workbook = lambda path, on_demand=False: wb
        sys.modules["xlrd"] = x
        try:
            (env.tmp / "lama.xls").write_bytes(b"\xd0\xcf\x11\xe0 palsu")
            s = main.read_sheets(env.tmp / "lama.xls")["Lembar1"]
        finally:
            del sys.modules["xlrd"]
        assert s.cols == ["Nama", "Usia", "Lahir", "Aktif", "Err"]
        assert s.rows == [["Budi", "30", "2024-01-15 00:00:00", "True", "#DIV/0!"], ["", "", "3.5", "False", "x"]]

    def test_deteksi_format_dari_isi_bukan_nama(self, env, sample):
        import shutil
        a = env.tmp / "dokumen"; shutil.copy(sample, a)
        b = make_ods(env.tmp / "ods.bin", {"S": ods_rows(SAMPLE)}); c = env.tmp / "lama"; c.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 20)
        d = env.tmp / "catatan"; d.write_text("hanya teks")
        assert [main.kind_of(x) for x in (a, b, c, d, env.tmp / "tidak-ada")] == ["xlsx", "ods", "xls", None, None]
        assert main.read_sheets(a)["Data"].rows[0][0] == "Budi Santoso"  # xlsx tanpa ekstensi tetap terbaca

    @pytest.mark.parametrize("penutup", ["/>", " />"])  # openpyxl menulis salah satunya, tergantung ada/tidaknya lxml
    def test_xlsx_dengan_dimensi_salah_tetap_terbaca_penuh(self, env, sample, penutup):
        out = env.tmp / "dimensi.xlsx"
        with zipfile.ZipFile(sample) as zi, zipfile.ZipFile(out, "w") as zo:
            n = 0
            for item in zi.infolist():
                data = zi.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data, n = re.subn(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1"' + penutup.encode(), data)
                zo.writestr(item, data)
        assert n == 1                                        # tag dimensi benar-benar diganti
        assert len(main.read_sheets(out)["Data"].rows) == 4  # tanpa reset_dimensions hanya sel A1 yang terbaca


# --------------------------------------------- 8. Android: penyimpanan & pemilih file
class TestAndroid:
    def test_adapter_pemilih_file_saf(self, env):
        calls = {}

        class Chooser:
            def __init__(self, cb): calls["cb"] = cb
            def choose_content(self, mime, multiple=False): calls["args"] = (mime, multiple)

        class SharedStorage:
            def copy_from_shared(self, uri):
                if uri == "gagal":
                    raise OSError("tidak bisa disalin")
                return None if uri == "kosong" else "/cache/" + uri.split("/")[-1]

        m = types.ModuleType("androidstorage4kivy"); m.Chooser, m.SharedStorage = Chooser, SharedStorage
        sys.modules["androidstorage4kivy"] = m
        try:
            got = []
            keep = main.android_pick(lambda paths, failed: got.append((paths, failed)))
            assert isinstance(keep, Chooser) and calls["args"] == ("*/*", True)   # semua jenis, boleh banyak
            calls["cb"](["content://x/a.xlsx", "gagal", "kosong", "content://x/b.csv"])
            assert got == []                                                      # belum: menunggu thread Kivy
            sys.modules["kivy.clock"].Clock.events[-1].fn(0)
            assert got == [(["/cache/a.xlsx", "/cache/b.csv"], 2)]
            calls["cb"](None); sys.modules["kivy.clock"].Clock.events[-1].fn(0)   # batal memilih
            assert got[-1] == ([], 0)
        finally:
            del sys.modules["androidstorage4kivy"]

    def test_tombol_tambah_file_di_android_memakai_pemilih_sistem(self, env, sample, monkeypatch):
        monkeypatch.setattr(main, "platform", "android")
        monkeypatch.setattr(main, "android_pick", lambda cb: (cb([sample, str(env.tmp / "hilang.xlsx")], 1), "chooser")[1])
        u = upl(); u.open_chooser()
        assert u.queue == [sample] and u._chooser == "chooser"
        assert env.pops[-1].title == "Sebagian file dilewati" and "[b]2[/b]" in env.pops[-1].msg

    def test_pemilih_android_gagal_dimuat_dilaporkan(self, env, monkeypatch):
        monkeypatch.setattr(main, "platform", "android")
        def boom(cb): raise ImportError("No module named 'jnius' [x]")
        monkeypatch.setattr(main, "android_pick", boom)
        upl().open_chooser()
        assert env.pops[-1].title == "Penyimpanan tidak dapat dibuka" and "jnius" in env.pops[-1].msg
        assert valid_markup(env.pops[-1].msg)

    def test_file_tanpa_ekstensi_disimpan_dengan_ekstensi_dan_terbaca(self, env, sample):
        import shutil
        raw = env.tmp / "dokumen-drive"; shutil.copy(sample, raw)       # nama dari Google Drive: tanpa ekstensi
        u = upl(); assert u.add_files([str(raw)]) == (1, 0)
        u.ask_save(); env.pops[-1].btn["Ya, simpan"]()
        item = main.load_index()[0]
        assert item["name"] == "dokumen-drive" and item["stored"].endswith(".xlsx")
        main.CACHE.clear(); d = dash(); d.reload()
        assert len(d.ids.rv.data) == 4

    def test_init_storage_membuat_folder_dan_data_bertahan(self, env, sample):
        base = env.tmp / "app_private" / "storage"
        main.init_storage(base)
        assert main.FILES_DIR.is_dir() and main.INDEX == base / "index.json"
        u = upl(); u.add_files([sample]); u.ask_save(); env.pops[-1].btn["Ya, simpan"]()
        stored = main.load_index()[0]["stored"]
        assert (base / "files" / stored).exists() and (base / "index.json").exists()
        main.CACHE.clear(); d = dash(); d.reload()                          # simulasi buka-ulang aplikasi
        assert len(d.ids.rv.data) == 4

    def test_data_disimpan_di_user_data_dir_bukan_di_samping_kode(self):
        assert "init_storage(Path(self.user_data_dir) / \"storage\")" in SOURCE

    def test_tombol_kembali_dan_jeda_aplikasi(self, env, sample):
        app = main.ExcelVaultApp(); app.root = SimpleNamespace(current="upload")
        assert app._on_key(None, 27) is False                              # belum ada data: keluar seperti biasa
        main.save_index([{"name": "a", "stored": "b"}])
        assert app._on_key(None, 27) is True and app.root.current == "dashboard"
        assert app._on_key(None, 27) is False                              # di dashboard: keluar
        app.root.current = "upload"; assert app._on_key(None, 97) is False  # tombol lain diabaikan
        assert app.on_pause() is True                                      # tanpa ini app mati saat pemilih file terbuka


# ------------------------------------------------------ 9. pengemasan (spec, workflow)
HERE = Path(__file__).resolve().parent


class TestPackaging:
    def spec(self):
        cp = configparser.ConfigParser(interpolation=None)
        cp.read(HERE / "buildozer.spec", encoding="utf-8")
        return cp["app"]

    def test_buildozer_spec(self):
        a = self.spec()
        reqs = {r.strip().split("==")[0] for r in a["requirements"].split(",")}
        assert {"python3", "kivy", "openpyxl", "et_xmlfile", "xlrd", "androidstorage4kivy", "pyjnius", "android"} <= reqs
        assert not ({"pandas", "numpy"} & reqs)                              # paket C: rawan gagal di Android
        assert re.fullmatch(r"[a-z][a-z0-9]*", a["package.name"]) and (HERE / "main.py").exists()
        assert a["android.accept_sdk_license"].lower() == "true"             # tanpa ini build CI menunggu input
        assert a["android.release_artifact"] == "apk" and int(a["android.minapi"]) >= 21
        assert "MANAGE_EXTERNAL_STORAGE" not in a.get("android.permissions", "")  # tidak perlu izin storage
        assert "test_*.py" in a["source.exclude_patterns"] and "py" in a["source.include_exts"]

    def test_semua_impor_pihak_ketiga_ada_di_requirements(self):
        names = set()
        for n in ast.walk(ast.parse(SOURCE)):
            if isinstance(n, ast.Import): names |= {x.name.split(".")[0] for x in n.names}
            if isinstance(n, ast.ImportFrom) and n.level == 0: names.add(n.module.split(".")[0])
        third = {n for n in names if n not in sys.stdlib_module_names}
        reqs = {r.strip().split("==")[0] for r in self.spec()["requirements"].split(",")}
        assert third and third <= reqs, third - reqs

    def test_workflow_github_actions(self):
        try:
            import yaml
        except ImportError:
            return
        wf = yaml.safe_load((HERE / ".github" / "workflows" / "build-apk.yml").read_text(encoding="utf-8"))
        on = wf.get("on", wf.get(True))
        assert "v*" in on["push"]["tags"] and "workflow_dispatch" in on
        assert wf["permissions"]["contents"] == "write"
        assert set(wf["jobs"]) == {"test", "build"} and wf["jobs"]["build"]["needs"] == "test"
        assert "refs/tags/v" in wf["jobs"]["build"]["if"] and "workflow_dispatch" in wf["jobs"]["build"]["if"]
        steps = wf["jobs"]["build"]["steps"] + wf["jobs"]["test"]["steps"]
        uses = [x["uses"] for x in steps if "uses" in x]
        assert all("@" in u for u in uses)                                   # semua action dikunci versinya
        for need in ("actions/checkout", "actions/setup-python", "actions/setup-java", "actions/cache",
                     "actions/upload-artifact", "softprops/action-gh-release"):
            assert any(u.startswith(need) for u in uses), need
        py = [x["with"]["python-version"] for x in steps if x.get("uses", "").startswith("actions/setup-python")]
        assert all(isinstance(v, str) for v in py)                           # '3.11' bukan angka 3.11
        text = (HERE / ".github" / "workflows" / "build-apk.yml").read_text(encoding="utf-8")
        for token in ("buildozer -v android debug", "buildozer -v android release", "P4A_RELEASE_KEYSTORE",
                      "P4A_RELEASE_KEYALIAS_PASSWD", "libffi-dev", "libtinfo6", "pytest -v test_main.py"):
            assert token in text, token

    def test_readme_memuat_alur_instalasi(self):
        r = (HERE / "README.md").read_text(encoding="utf-8")
        for token in ("Releases", "Actions", "KEYSTORE_BASE64", "Install unknown apps", "v1.0.0"):
            assert token in r, token


# ---------------------------------------------------------------- 10. kinerja
class TestPerformance:
    def test_cari_400rb_sel_dan_bangun_3000_kartu(self):
        rng = random.Random(1)
        pool = ["Budi Santoso", "Ani", "Jakarta", "Bandung", "kopi", "teh", "12345", "Citra Dewi"]
        rows = [[rng.choice(pool) for _ in range(8)] for _ in range(50_000)]
        t = [{"file": "f", "sheet": "s", "cols": list("ABCDEFGH"), "rows": rows, "first_row": 2}]
        t0 = time.perf_counter(); r = main.search_tables(t, "budi"); t1 = time.perf_counter()
        cards = [main.hit_card(t[0], rr, c, "budi") for _, rr, c in r["hits"]]; t2 = time.perf_counter()
        print("\n[kinerja] cari 400rb sel: %.2fs | 3000 kartu: %.2fs | hit=%d" % (t1 - t0, t2 - t1, r["cells"]))
        assert r["cells"] > 40_000 and len(cards) == main.MAX_RESULTS
        assert t1 - t0 < 3 and t2 - t1 < 3

    def test_baca_xlsx_20rb_baris(self, env):
        p = env.tmp / "besar.xlsx"
        wb = openpyxl.Workbook(write_only=True); ws = wb.create_sheet("S")
        ws.append(["Nama", "Kota", "Usia", "Catatan", "Kode", "Alamat"])
        for i in range(20_000):
            ws.append(["Nama %d" % i, "Kota %d" % (i % 50), i % 90, "catatan", "K%05d" % i, "Jl. Contoh %d" % i])
        wb.save(p)
        t0 = time.perf_counter(); s = main.read_sheets(p)["S"]; t1 = time.perf_counter()
        r = main.search_tables([{"file": "b", "sheet": "S", "cols": s.cols, "rows": s.rows, "first_row": s.first_row}], "kota 7")
        print("\n[kinerja] baca 20rb baris xlsx: %.2fs | cari: %.3fs" % (t1 - t0, time.perf_counter() - t1))
        assert len(s.rows) == 20_000 and r["cells"] == 400 and t1 - t0 < 20
